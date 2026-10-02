"""Edge-case tests. Run with:  python -m pytest -q"""
import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from peerlens.excel_export import ExcelDashboard
from peerlens.live import standardise_statements
from peerlens.pipeline import run_analysis
from peerlens.ratios import RATIO_BY_KEY, RatioEngine, safe_divide
from peerlens.synthetic import generate_financials
from peerlens.universe import load_nifty500
from peerlens.validator import FinancialDataValidator

UNIVERSE, _ = load_nifty500(allow_network=False)
CEMENT = UNIVERSE[UNIVERSE.industry == "Construction Materials"].head(5)[["company", "symbol"]]


def test_bundled_universe_is_a_real_nifty_500():
    assert len(UNIVERSE) >= 490
    assert {"company", "industry", "symbol", "yahoo_symbol"} <= set(UNIVERSE.columns)
    assert UNIVERSE.yahoo_symbol.str.endswith(".NS").all()


def test_safe_divide_handles_zero_and_missing():
    result = safe_divide([10, 5, 1], [0, np.nan, 4])
    assert pd.isna(result[0]) and pd.isna(result[1]) and result[2] == 0.25


def test_generator_is_reproducible_and_balances():
    a = generate_financials(CEMENT, "Construction Materials", seed=7)
    b = generate_financials(CEMENT, "Construction Materials", seed=7)
    pd.testing.assert_frame_equal(a, b)
    gap = a.total_assets - (a.equity + a.total_debt + a.current_liabilities)
    assert gap.abs().max() <= 1.0


def test_validator_finds_and_fixes_planted_errors():
    raw = generate_financials(CEMENT, "Construction Materials", seed=42, inject_errors=True)
    validator = FinancialDataValidator(mode="synthetic")
    before = validator.validate(raw)
    assert {"Duplicate row", "Negative value", "Missing value"} <= set(before["check"])
    clean, fixes = validator.clean(raw)
    assert validator.validate(clean).empty
    assert "Rebuilt from identity" in set(fixes["action"])


def test_missing_column_is_rejected_clearly():
    raw = generate_financials(CEMENT, "Construction Materials")
    with pytest.raises(ValueError, match="revenue"):
        FinancialDataValidator().validate(raw.drop(columns=["revenue"]))


def test_rating_both_directions_and_na():
    assert RATIO_BY_KEY["roce"].rate(0.20) == "Green"
    assert RATIO_BY_KEY["roce"].rate(0.05) == "Red"
    assert RATIO_BY_KEY["debt_to_equity"].rate(0.2) == "Green"
    assert RATIO_BY_KEY["debt_to_equity"].rate(1.6) == "Red"
    assert RATIO_BY_KEY["interest_coverage"].rate(np.nan) == "n/a"


def test_empty_input_does_not_crash():
    raw = generate_financials(CEMENT, "Construction Materials")
    assert RatioEngine().compute_all(raw.iloc[0:0]).empty


def _fake_yahoo(bank=False):
    """Statements shaped like yfinance output: rows are line items, columns are period-end dates."""
    dates = pd.to_datetime(["2025-03-31", "2024-03-31", "2023-03-31"])
    income = pd.DataFrame({d: {"Total Revenue": 7.0e11 - i * 5e10, "EBITDA": 1.4e11, "EBIT": 1.0e11,
                               "Reconciled Depreciation": 4e10, "Interest Expense": 1e10,
                               "Pretax Income": 9e10, "Tax Provision": 2.2e10, "Net Income": 6.8e10}
                           for i, d in enumerate(dates)})
    balance_items = {"Total Assets": 1.2e12, "Stockholders Equity": 6e11, "Total Debt": 1.5e11,
                     "Cash And Cash Equivalents": 3e10, "Accounts Receivable": 4e10}
    if not bank:
        balance_items.update({"Current Assets": 2.5e11, "Current Liabilities": 2.2e11, "Inventory": 6e10})
    balance = pd.DataFrame({d: balance_items for d in dates})
    return income, balance


def test_yahoo_statements_are_standardised_to_crore_and_fiscal_years():
    income, balance = _fake_yahoo()
    data = standardise_statements(income, balance, "Test Co", "TEST")
    assert list(data.year) == ["FY2023", "FY2024", "FY2025"]
    assert data.loc[data.year == "FY2025", "revenue"].iloc[0] == pytest.approx(70000.0)


def test_live_mode_bank_has_no_liquidity_ratios_but_still_runs():
    income, balance = _fake_yahoo(bank=True)
    frames = [standardise_statements(income * f, balance * f, f"Bank {i}", f"B{i}")
              for i, f in enumerate([1.0, 0.8, 1.3])]
    result = run_analysis(pd.concat(frames, ignore_index=True), "Bank 0", mode="live")
    keys = [r.key for r in result.ratios]
    assert "current_ratio" not in keys and "roe" in keys
    assert result.rank >= 1


def test_excel_export_has_formulas_and_matching_layout():
    raw = generate_financials(CEMENT, "Construction Materials", seed=42, inject_errors=True)
    result = run_analysis(raw, CEMENT.company.iloc[0], mode="synthetic")
    payload = ExcelDashboard(result.clean, result.ratios, result.focal, result.year, "Construction Materials",
                             "test", result.issues_before, result.fixes, result.recommendations,
                             result.summary_line).to_bytes()
    wb = load_workbook(io.BytesIO(payload))
    assert {"Dashboard", "Data", "Ratios", "Thresholds", "Data Quality", "Notes"} <= set(wb.sheetnames)
    assert str(wb["Ratios"]["D2"].value).startswith("=IFERROR(")


def _cement_run(focal):
    names = ["J.K. Cement", "Dalmia Bharat", "UltraTech Cement", "India Cements", "Nuvoco Vistas Corporation"]
    peers = UNIVERSE[UNIVERSE.company.isin(names)][["company", "symbol"]]
    return run_analysis(generate_financials(peers, "Construction Materials", seed=42, inject_errors=True),
                        focal, mode="synthetic")


def test_rank_counts_only_strictly_higher_scores():
    """Rank = 1 + number of peers with a strictly higher score, so tied companies share a rank."""
    names = ["J.K. Cement", "Dalmia Bharat", "UltraTech Cement", "India Cements", "Nuvoco Vistas Corporation"]
    for focal in names:
        result = _cement_run(focal)
        scores = result.health[result.health.year == result.year].set_index("company")["health_score"]
        expected = 1 + int((scores > scores[focal]).sum())
        assert result.rank == expected
        assert result.tied == bool((scores == scores[focal]).sum() > 1)


def test_a_ratio_is_never_both_a_strength_and_a_finding():
    for company in ["J.K. Cement", "Dalmia Bharat", "UltraTech Cement"]:
        result = _cement_run(company)
        flagged = {n["ratio"] for n in result.recommendations}
        assert not any(s.startswith(tuple(flagged)) for s in result.strengths)


def test_benchmark_excludes_the_company_itself():
    result = _cement_run("J.K. Cement")
    snap = result.ratio_values[result.ratio_values.year == result.year]
    others = snap[snap.company != "J.K. Cement"]["current_ratio"].median()
    note = next(n for n in result.recommendations if n["ratio"] == "Current ratio")
    assert RATIO_BY_KEY["current_ratio"].format(others) in note["text"]


# ---------------------------------------------------------------- macro module
from peerlens import macro as macro_module  # noqa: E402
from peerlens.synthetic import YEARS, true_gdp_beta  # noqa: E402


def test_macro_path_is_reproducible_and_follows_its_rules():
    a, b = macro_module.generate_macro(42), macro_module.generate_macro(42)
    pd.testing.assert_frame_equal(a, b)
    assert a.loc[a.year == "FY2021", "gdp_growth"].iloc[0] < 0          # pandemic contraction
    assert a.repo_rate.between(0.035, 0.08).all()                          # policy rate stays in its band
    assert (a.cpi_inflation >= 0.02).all()
    # Taylor-style rule: higher inflation means a higher policy rate
    assert np.corrcoef(a.cpi_inflation, a.repo_rate)[0, 1] > 0.5


def test_regression_recovers_the_built_in_cyclicality():
    path = macro_module.macro_for_years(YEARS)
    for industry in ["Metals & Mining", "Construction Materials", "Fast Moving Consumer Goods", "Healthcare"]:
        peers = UNIVERSE[UNIVERSE.industry == industry][["company", "symbol"]]
        panel = macro_module.growth_panel(generate_financials(peers, industry, seed=42), path)
        estimate = macro_module.pooled_sensitivity(panel)["gdp_beta"]
        assert abs(estimate - true_gdp_beta(industry)) < 0.25, industry


def test_cyclical_industries_rank_above_defensive_ones():
    path = macro_module.macro_for_years(YEARS)
    betas = {}
    for industry in ["Metals & Mining", "Healthcare"]:
        peers = UNIVERSE[UNIVERSE.industry == industry][["company", "symbol"]]
        panel = macro_module.growth_panel(generate_financials(peers, industry, seed=42), path)
        betas[industry] = macro_module.pooled_sensitivity(panel)
    assert betas["Metals & Mining"]["profile"] == "Highly cyclical"
    assert betas["Healthcare"]["profile"] == "Defensive"


def test_short_history_gives_no_estimate_instead_of_a_fake_one():
    peers = CEMENT[["company", "symbol"]]
    short = generate_financials(peers, "Construction Materials", seed=42).query("year >= 'FY2023'")
    panel = macro_module.growth_panel(short, macro_module.macro_for_years(YEARS))
    sensitivity = macro_module.gdp_sensitivity(panel)
    assert sensitivity["beta"].isna().all()        # 2 growth points per company: refuse to estimate


def test_generator_keeps_every_company_solvent():
    for industry in ["Power", "Telecommunication", "Financial Services", "Realty"]:
        peers = UNIVERSE[UNIVERSE.industry == industry][["company", "symbol"]]
        data = generate_financials(peers, industry, seed=7)
        assert (data.equity > 0).all(), industry
